"""Extract plain text from an uploaded resume.

Text is extracted once at upload and stored, because every later call (skill
extraction, match scoring, document generation) needs the resume and re-parsing
a PDF on each one would be wasted work.

Local parsing first: pypdf and python-docx cost nothing. Only when a PDF yields
too little text - a scan, or a design-heavy CV exported as images - do we fall
back to the model reading the PDF natively, which costs a call but is the
only thing that works on an image-only document.
"""

from __future__ import annotations

import io
import logging

from app.config import settings

logger = logging.getLogger(__name__)

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TEXT_TYPES = {"text/plain"}
ALLOWED_RESUME_TYPES = PDF_TYPES | DOCX_TYPES | TEXT_TYPES

# Below this, local parsing has effectively failed - a real resume has more.
MIN_USABLE_CHARS = 200


class ResumeParseError(ValueError):
    """The upload could not be read - message is safe to show the user."""


def _from_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Local PDF parse failed: %s", exc)
        return ""


def _from_docx(raw: bytes) -> str:
    try:
        import docx

        document = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs]
        # Resumes frequently lay out dates and roles in tables, which are not
        # in `paragraphs` at all - miss these and half the CV disappears.
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(p for p in parts if p and p.strip()).strip()
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError("That .docx file could not be read.") from exc


def _from_pdf_via_model(raw: bytes) -> str:
    """Last resort for scanned PDFs. Costs one API call.

    Gemini reads PDFs natively, so a scan that pypdf cannot touch still works.
    It walks the same model chain as every other AI call, so one model being
    out of quota does not break uploads.
    """
    if not settings.ai_enabled:
        return ""
    from google.genai import types

    from app.services import ai

    logger.info("Falling back to the model for PDF text extraction (likely a scan)")
    contents = [
        types.Part.from_bytes(data=raw, mime_type="application/pdf"),
        "Transcribe this resume as plain text, preserving section headings and "
        "bullet points. Output only the transcription.",
    ]
    for model in ai.models_to_try():
        try:
            response = ai._client().models.generate_content(model=model, contents=contents)
        except Exception as exc:  # noqa: BLE001
            if not ai.remember_quota_error(model, exc):
                logger.warning("PDF transcription failed on %s: %s", model, str(exc)[:200])
            continue
        return (response.text or "").strip()
    # Every model failed; the caller reports the file as unreadable.
    return ""


def extract(raw: bytes, content_type: str) -> str:
    """Return the resume's plain text, or raise ResumeParseError."""
    if len(raw) > settings.max_resume_bytes:
        limit_mb = settings.max_resume_bytes // (1024 * 1024)
        raise ResumeParseError(f"That file is too large (max {limit_mb} MB).")

    if content_type in DOCX_TYPES:
        text = _from_docx(raw)
    elif content_type in TEXT_TYPES:
        text = raw.decode("utf-8", errors="replace").strip()
    elif content_type in PDF_TYPES:
        text = _from_pdf(raw)
        if len(text) < MIN_USABLE_CHARS:
            text = _from_pdf_via_model(raw) or text
    else:
        raise ResumeParseError("Upload a PDF, DOCX or plain text resume.")

    if len(text) < MIN_USABLE_CHARS:
        raise ResumeParseError(
            "Could not read enough text from that file. If it is a scanned image, "
            "try exporting a text-based PDF."
        )
    return text
